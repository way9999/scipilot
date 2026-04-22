# -*- coding: utf-8 -*-
"""完整项目功能流程测试"""
import json
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
sys.stdout.reconfigure(encoding='utf-8')

PASS, FAIL, SKIP = 0, 0, 0
SKIP_NETWORK_TESTS = os.environ.get("SCIPILOT_SKIP_NETWORK_TESTS", "").strip().lower() in {"1", "true", "yes", "on"}

def test(name, fn):
    global PASS, FAIL, SKIP
    try:
        result = fn()
        if result == "SKIP":
            SKIP += 1
            print(f"  SKIP  {name}")
        else:
            PASS += 1
            print(f"  PASS  {name}")
    except Exception as e:
        FAIL += 1
        print(f"  FAIL  {name}: {e}")


def network_test(name, fn):
    if SKIP_NETWORK_TESTS:
        return test(name, lambda: "SKIP")
    return test(name, fn)

# ============================================================
print("\n=== 1. 文件结构完整性 ===")
# ============================================================
required_files = [
    ".claude/commands/research.md",
    ".claude/commands/focus.md",
    ".claude/commands/lit-verify.md",
    ".claude/commands/lit-download.md",
    ".claude/commands/proposal.md",
    ".claude/commands/review-write.md",
    ".claude/commands/paper-write.md",
    "tools/crossref_search.py",
    "tools/semantic_scholar.py",
    "tools/zotero_import.py",
    "tools/arxiv_download.py",
    "tools/scholarly_search.py",
    "tools/paperscraper_tool.py",
    "tools/pypaperbot_tool.py",
    "tools/scihub2pdf_tool.py",
    "tools/unified_search.py",
    "tools/project_models.py",
    "tools/project_state.py",
    "tools/paper_dashboard.py",
    "tools/experiment_design.py",
    "tools/literature_review.py",
    "tools/research_capability_audit.py",
    "tools/research_export.py",
    "tools/research_qa.py",
    "tools/requirements.txt",
    "references/experiment-templates.md",
    "references/step1-intent/domain-adapters.md",
    "references/step2-factcheck/fact-check.md",
    "references/step3-structure/outline-template.md",
    "references/step4-generation/writing-style-adapters.md",
    "CLAUDE.md",
]
for f in required_files:
    test(f"文件存在: {f}", lambda f=f: None if os.path.exists(f) else (_ for _ in ()).throw(FileNotFoundError(f)))

# ============================================================
print("\n=== 2. 工具模块导入 ===")
# ============================================================
test("crossref_search", lambda: __import__("tools.crossref_search"))
test("semantic_scholar", lambda: __import__("tools.semantic_scholar"))
test("zotero_import", lambda: __import__("tools.zotero_import"))
test("arxiv_download", lambda: __import__("tools.arxiv_download"))
test("scholarly_search", lambda: __import__("tools.scholarly_search"))
test("paperscraper_tool", lambda: __import__("tools.paperscraper_tool"))
test("pypaperbot_tool", lambda: __import__("tools.pypaperbot_tool"))
test("scihub2pdf_tool", lambda: __import__("tools.scihub2pdf_tool"))
test("unified_search", lambda: __import__("tools.unified_search"))
test("project_models", lambda: __import__("tools.project_models"))
test("project_state", lambda: __import__("tools.project_state"))
test("paper_dashboard", lambda: __import__("tools.paper_dashboard"))
test("experiment_design", lambda: __import__("tools.experiment_design"))
test("literature_review", lambda: __import__("tools.literature_review"))
test("research_capability_audit", lambda: __import__("tools.research_capability_audit"))
test("research_export", lambda: __import__("tools.research_export"))
test("research_qa", lambda: __import__("tools.research_qa"))

# ============================================================
print("\n=== 3. 统一搜索路由逻辑 (本地) ===")
# ============================================================
from tools.unified_search import DISCIPLINE_ROUTES, DOWNLOAD_CHAIN

def test_routing():
    assert "cs" in DISCIPLINE_ROUTES
    assert "bio" in DISCIPLINE_ROUTES
    assert "chemistry" in DISCIPLINE_ROUTES
    assert "arxiv" in DISCIPLINE_ROUTES["cs"]["search"]
    assert "pubmed" in DISCIPLINE_ROUTES["bio"]["search"]
    assert len(DOWNLOAD_CHAIN) == 5
test("学科路由表完整", test_routing)

def test_title_match():
    from tools.unified_search import _title_match
    assert _title_match("Attention Is All You Need", "attention is all you need")
    assert not _title_match("Paper A", "Paper B")
test("title match logic", test_title_match)


def test_normalize_paper_dict():
    from tools.project_models import normalize_paper_dict

    paper = normalize_paper_dict({
        "title": "Attention Is All You Need",
        "authors": ["Ashish Vaswani"],
        "year": "2017",
        "journal": "NeurIPS",
        "doi": "10.5555/3295222.3295349",
    }, source="crossref", discipline="cs")

    assert paper["record_id"].startswith("doi:")
    assert paper["venue"] == "NeurIPS"
    assert paper["source"] == "crossref"
    assert paper["discipline"] == "cs"
test("paper record normalization", test_normalize_paper_dict)


def test_dedupe_and_merge_records():
    from tools.project_models import dedupe_papers

    papers = dedupe_papers([
        {
            "title": "Attention Is All You Need",
            "authors": ["Ashish Vaswani"],
            "year": 2017,
            "doi": "10.5555/3295222.3295349",
            "source": "crossref",
        },
        {
            "title": "Attention Is All You Need",
            "authors": ["A. Vaswani"],
            "year": 2017,
            "doi": "10.5555/3295222.3295349",
            "abstract": "Transformers replace recurrence.",
            "pdf_url": "https://example.com/paper.pdf",
            "source": "semantic_scholar",
        },
    ])

    assert len(papers) == 1
    assert papers[0]["abstract"] == "Transformers replace recurrence."
    assert "crossref" in papers[0]["sources"]
    assert "semantic_scholar" in papers[0]["sources"]
test("paper record dedupe", test_dedupe_and_merge_records)


def test_project_state_detection():
    import tempfile
    from pathlib import Path
    from tools.project_state import detect_project_state, register_search_results

    with tempfile.TemporaryDirectory() as tmpdir:
        root = Path(tmpdir)
        (root / "knowledge-base").mkdir()
        (root / "drafts").mkdir()
        (root / "output").mkdir()
        (root / "papers").mkdir()
        (root / "outline.md").write_text("---\nstatus: frozen\n---\n", encoding="utf-8")

        register_search_results([
            {
                "title": "Example Paper",
                "authors": ["Alice"],
                "year": 2024,
                "source": "crossref",
            }
        ], project_root=root, discipline="cs", query="example")

        state = detect_project_state(root)
        assert state["current_stage"] == "writing"
        assert state["outline_frozen"] is True
        assert state["artifacts"]["paper_index_count"] == 1
test("project state detection", test_project_state_detection)


def test_route_recommendation():
    import shutil
    import uuid
    from pathlib import Path
    from tools.project_state import recommend_next_route, register_search_results

    root = Path("output") / f"tmp-route-test-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        (root / "knowledge-base").mkdir()
        register_search_results([
            {"title": "Example Paper", "authors": ["Alice"], "year": 2024, "source": "crossref"}
        ], project_root=root, discipline="cs", query="example")
        route = recommend_next_route(arguments="download these PDFs", project_root=root)
        assert route["recommended_route"] == "/lit-download"
        assert route["current_stage"] == "literature"
    finally:
        shutil.rmtree(root)
test("route recommendation", test_route_recommendation)


def test_paper_dashboard_build():
    import shutil
    import uuid
    from pathlib import Path
    from tools.paper_dashboard import build_dashboard
    from tools.project_state import register_search_results

    root = Path("output") / f"tmp-dashboard-test-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        (root / "knowledge-base").mkdir()
        (root / "drafts").mkdir()
        (root / "output").mkdir()
        (root / "papers").mkdir()
        (root / "outline.md").write_text("---\nstatus: frozen\n---\n", encoding="utf-8")
        register_search_results([
            {
                "title": "Dashboard Paper",
                "authors": ["Alice", "Bob"],
                "year": 2025,
                "source": "crossref",
                "verified": True,
            }
        ], project_root=root, discipline="bio", query="dashboard")
        output_path = build_dashboard(project_root=root)
        html = output_path.read_text(encoding="utf-8")
        assert output_path.exists()
        assert "Paper Workbench" in html
        assert "Dashboard Paper" in html
    finally:
        shutil.rmtree(root)
test("paper dashboard build", test_paper_dashboard_build)


def test_project_runner_blocks_agent_edits_by_default():
    import shutil
    import types
    import uuid
    from pathlib import Path

    import tools.project_runner as project_runner

    root = Path("output") / f"tmp-project-runner-safe-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    (root / "main.py").write_text("print('hello from test')\n", encoding="utf-8")

    original_run_python = project_runner._run_python
    original_collect_output_data = project_runner._collect_output_data
    original_agent_bridge = sys.modules.get("tools.agent_bridge")
    original_allow_env = os.environ.pop("SCIPILOT_AGENT_ALLOW_PROJECT_EDITS", None)

    run_calls = {"count": 0}
    agent_calls = []

    try:
        def fake_run_python(project_root, entry_script, timeout=300):
            run_calls["count"] += 1
            return {"success": False, "stdout": "", "stderr": "boom", "returncode": 1}

        def fake_collect_output_data(project_root):
            return {"figures": [], "csv_files": [], "json_files": [], "log_lines": [], "metrics": {}}

        def fake_run_agent_task(**kwargs):
            agent_calls.append(kwargs)
            return {"success": True, "output": "fixed", "files_modified": ["main.py"]}

        fake_agent_bridge = types.ModuleType("tools.agent_bridge")
        fake_agent_bridge.agent_enabled = lambda: True
        fake_agent_bridge.run_agent_task = fake_run_agent_task
        fake_agent_bridge.results_sufficient = lambda collected, project_context=None: False
        fake_agent_bridge.build_fix_prompt = lambda run_result, project_type, entry_script=None: "fix"
        fake_agent_bridge.build_supplement_prompt = (
            lambda collected, project_type, project_context=None: "supplement"
        )
        fake_agent_bridge._agent_config = lambda: {"auto_fix": True, "auto_supplement": True}

        project_runner._run_python = fake_run_python
        project_runner._collect_output_data = fake_collect_output_data
        sys.modules["tools.agent_bridge"] = fake_agent_bridge

        result = project_runner.run_project(root)

        assert result["run_result"]["success"] is False
        assert run_calls["count"] == 1
        assert agent_calls == []
    finally:
        project_runner._run_python = original_run_python
        project_runner._collect_output_data = original_collect_output_data
        if original_agent_bridge is None:
            sys.modules.pop("tools.agent_bridge", None)
        else:
            sys.modules["tools.agent_bridge"] = original_agent_bridge
        if original_allow_env is None:
            os.environ.pop("SCIPILOT_AGENT_ALLOW_PROJECT_EDITS", None)
        else:
            os.environ["SCIPILOT_AGENT_ALLOW_PROJECT_EDITS"] = original_allow_env
        shutil.rmtree(root)
test("project_runner blocks agent edits by default", test_project_runner_blocks_agent_edits_by_default)


def test_project_runner_allows_opt_in_agent_edits():
    import shutil
    import types
    import uuid
    from pathlib import Path

    import tools.project_runner as project_runner

    root = Path("output") / f"tmp-project-runner-opt-in-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    (root / "main.py").write_text("print('hello from test')\n", encoding="utf-8")

    original_run_python = project_runner._run_python
    original_collect_output_data = project_runner._collect_output_data
    original_agent_bridge = sys.modules.get("tools.agent_bridge")
    original_allow_env = os.environ.pop("SCIPILOT_AGENT_ALLOW_PROJECT_EDITS", None)

    responses = [
        {"success": False, "stdout": "", "stderr": "boom", "returncode": 1},
        {"success": True, "stdout": "ok", "stderr": "", "returncode": 0},
    ]
    run_calls = {"count": 0}
    agent_calls = []

    try:
        def fake_run_python(project_root, entry_script, timeout=300):
            run_calls["count"] += 1
            return responses.pop(0)

        def fake_collect_output_data(project_root):
            return {"figures": [], "csv_files": [], "json_files": [], "log_lines": [], "metrics": {}}

        def fake_run_agent_task(**kwargs):
            agent_calls.append(kwargs)
            return {"success": True, "output": "fixed", "files_modified": ["main.py"]}

        fake_agent_bridge = types.ModuleType("tools.agent_bridge")
        fake_agent_bridge.agent_enabled = lambda: True
        fake_agent_bridge.run_agent_task = fake_run_agent_task
        fake_agent_bridge.results_sufficient = lambda collected, project_context=None: True
        fake_agent_bridge.build_fix_prompt = lambda run_result, project_type, entry_script=None: "fix"
        fake_agent_bridge.build_supplement_prompt = (
            lambda collected, project_type, project_context=None: "supplement"
        )
        fake_agent_bridge._agent_config = lambda: {"auto_fix": True, "auto_supplement": False}

        project_runner._run_python = fake_run_python
        project_runner._collect_output_data = fake_collect_output_data
        sys.modules["tools.agent_bridge"] = fake_agent_bridge

        result = project_runner.run_project(root, allow_agent_modifications=True)

        assert result["run_result"]["success"] is True
        assert run_calls["count"] == 2
        assert len(agent_calls) == 1
        assert agent_calls[0]["task"] == "fix"
    finally:
        project_runner._run_python = original_run_python
        project_runner._collect_output_data = original_collect_output_data
        if original_agent_bridge is None:
            sys.modules.pop("tools.agent_bridge", None)
        else:
            sys.modules["tools.agent_bridge"] = original_agent_bridge
        if original_allow_env is None:
            os.environ.pop("SCIPILOT_AGENT_ALLOW_PROJECT_EDITS", None)
        else:
            os.environ["SCIPILOT_AGENT_ALLOW_PROJECT_EDITS"] = original_allow_env
        shutil.rmtree(root)
test("project_runner allows explicit agent edit opt-in", test_project_runner_allows_opt_in_agent_edits)

# ============================================================
print("\n=== 4. 实验设计工具 (本地) ===")
# ============================================================
from tools.experiment_design import (
    ablation_study, baseline_comparison, hyperparameter_grid,
    hyperparameter_lhs, dose_response, screening_plate,
    clinical_groups, factorial_design,
)

def test_ablation():
    exps = ablation_study({
        "attention": ["multi-head", "single-head", "none"],
        "dropout": [0.1, 0.0],
    })
    assert len(exps) == 4  # full + 2 attention variants + 1 dropout variant
    assert exps[0]["name"] == "Full Model"
test("消融实验设计", test_ablation)

def test_baseline():
    m = baseline_comparison("Ours", ["BERT", "GPT"], ["GLUE", "SQuAD"], ["Acc", "F1"])
    assert m["total_runs"] == 6
    assert "table_template" in m
    assert "| Ours |" in m["table_template"]
test("基线对比矩阵", test_baseline)

def test_grid():
    configs = hyperparameter_grid({"lr": [1e-3, 1e-4], "bs": [16, 32]})
    assert len(configs) == 4
test("网格搜索", test_grid)

def test_lhs():
    samples = hyperparameter_lhs({"lr": (1e-5, 1e-2), "wd": (0, 0.1)}, n_samples=10)
    assert len(samples) == 10
    assert all(1e-5 <= s["lr"] <= 1e-2 for s in samples)
test("拉丁超立方采样", test_lhs)

def test_dose():
    groups = dose_response([0.1, 1, 10, 100], replicates=3)
    assert len(groups) == 5  # control + 4 doses
    assert groups[0]["group"] == "Vehicle Control"
test("剂量-反应设计", test_dose)

def test_plate():
    p = screening_plate(["A", "B", "C"], [0.1, 1, 10], replicates=2, plate_size=96)
    assert p["compounds"] == 3
    assert p["total_wells"] == 18
    assert p["plates_needed"] >= 1
test("筛选板布局", test_plate)

def test_clinical():
    d = clinical_groups(["Vehicle", "Low", "High"], n_per_arm=8, stratify_by=["weight"])
    assert d["total_n"] == 24
    assert "weight" in d["stratification"]
test("临床分组设计", test_clinical)

def test_factorial():
    exps = factorial_design({"A": [1, 2], "B": ["x", "y"], "C": [True, False]}, "full")
    assert len(exps) == 8
test("全因子设计", test_factorial)

# ============================================================
print("\n=== 5. Optuna 集成 (本地) ===")
# ============================================================
from tools.experiment_design import (
    create_optuna_study, optuna_search_template, get_optuna_results,
)

def test_optuna_study():
    study = create_optuna_study("test_study", direction="minimize")
    study.optimize(lambda trial: trial.suggest_float("x", -10, 10) ** 2, n_trials=10)
    r = get_optuna_results(study)
    assert r["n_trials"] == 10
    assert abs(r["best_params"]["x"]) < 5  # should find near 0
test("Optuna 超参搜索", test_optuna_study)

def test_optuna_template():
    code = optuna_search_template(
        {"lr": ("log_float", 1e-5, 1e-1), "hidden": ("int", 64, 512)},
        n_trials=20,
    )
    assert "import optuna" in code
    assert "suggest_float" in code
    assert "suggest_int" in code
test("Optuna 代码模板生成", test_optuna_template)

# ============================================================
print("\n=== 6. TDC 集成 ===")
# ============================================================
from tools.experiment_design import list_tdc_benchmarks, tdc_experiment_template

def test_tdc_benchmarks():
    b = list_tdc_benchmarks()
    assert "admet" in b
    assert "dti" in b
    assert "DAVIS" in b["dti"]["tasks"]
test("TDC 基准列表", test_tdc_benchmarks)

def test_tdc_template():
    code = tdc_experiment_template("ADME", "Caco2_Wang")
    assert "from tdc" in code
    assert "Caco2_Wang" in code
test("TDC 代码模板生成", test_tdc_template)

def test_tdc_load():
    from tools.experiment_design import load_tdc_dataset
    data = load_tdc_dataset("ADME", "Caco2_Wang", path="data/")
    assert data["info"]["train_size"] > 0
    assert data["info"]["test_size"] > 0
    print(f"         (train={data['info']['train_size']}, valid={data['info']['valid_size']}, test={data['info']['test_size']})")
test("TDC 数据集加载 (Caco2_Wang)", test_tdc_load)

# ============================================================
print("\n=== 7. 数据模型与桥接 (本地) ===")
# ============================================================
from tools.project_models import (
    PaperRecord, paper_identity, paper_quality_score,
    merge_paper_dicts, normalize_source_name, build_record_id,
)

def test_paper_record_creation():
    rec = PaperRecord(
        title="Test Paper",
        authors=["Alice", "Bob"],
        year="2024",
        doi="10.1234/test",
        source="crossref",
    )
    assert rec.record_id == "doi:10.1234/test"
    assert rec.year == 2024
    assert rec.source == "crossref"
    assert "crossref" in rec.sources
test("PaperRecord dataclass", test_paper_record_creation)


def test_paper_record_from_raw():
    raw = {
        "title": "Raw Paper",
        "authors": ["Charlie"],
        "year": 2023,
        "journal": "Nature",
        "doi": "10.5678/raw",
        "extra_field": "should go to metadata",
    }
    rec = PaperRecord.from_raw(raw, source="semantic_scholar", discipline="bio")
    assert rec.venue == "Nature"
    assert rec.discipline == "bio"
    assert rec.metadata.get("extra_field") == "should go to metadata"
test("PaperRecord.from_raw", test_paper_record_from_raw)


def test_paper_identity():
    assert paper_identity({"doi": "10.1234/x"}) == "doi:10.1234/x"
    assert paper_identity({"arxiv_id": "2301.00001"}) == "arxiv:2301.00001"
    assert paper_identity({"title": "My Paper", "year": 2024}).startswith("title:")
test("paper_identity", test_paper_identity)


def test_paper_quality_score():
    low = paper_quality_score({"title": "A"})
    high = paper_quality_score({
        "title": "B",
        "abstract": "text",
        "doi": "10/x",
        "venue": "ICML",
        "verified": True,
        "downloaded": True,
        "source": "crossref",
        "citation_count": 50,
    })
    assert high > low
test("paper_quality_score", test_paper_quality_score)


def test_merge_paper_dicts():
    a = {"title": "Paper X", "doi": "10.1/x", "source": "crossref", "year": 2024}
    b = {"title": "Paper X", "doi": "10.1/x", "source": "semantic_scholar",
         "year": 2024, "abstract": "Great paper", "pdf_url": "http://example.com/x.pdf"}
    merged = merge_paper_dicts(a, b)
    assert merged["abstract"] == "Great paper"
    assert "crossref" in merged["sources"]
    assert "semantic_scholar" in merged["sources"]
test("merge_paper_dicts", test_merge_paper_dicts)


def test_normalize_source_name():
    assert normalize_source_name("scholarly") == "google_scholar"
    assert normalize_source_name("CrossRef") == "crossref"
    assert normalize_source_name(None) == ""
test("normalize_source_name", test_normalize_source_name)


def test_build_record_id_priority():
    # DOI takes priority over arxiv_id
    rid = build_record_id("Title", year=2024, doi="10.1/x", arxiv_id="2301.00001")
    assert rid == "doi:10.1/x"
    # arxiv_id takes priority over title
    rid2 = build_record_id("Title", year=2024, arxiv_id="2301.00001")
    assert rid2 == "arxiv:2301.00001"
    # Fallback to title
    rid3 = build_record_id("My Title", year=2024)
    assert rid3.startswith("title:")
test("build_record_id 优先级", test_build_record_id_priority)


def test_research_bridge_parser():
    from tools.research_bridge import _build_parser
    parser = _build_parser()
    args = parser.parse_args(["search", "--query", "test", "--discipline", "cs"])
    assert args.query == "test"
    assert args.discipline == "cs"
    assert args.limit == 10
    assert not args.download
    args2 = parser.parse_args(["download", "--record-id", "doi:10/x"])
    assert args2.record_id == "doi:10/x"
    args3 = parser.parse_args(["refresh"])
    assert args3.command == "refresh"
    args4 = parser.parse_args(["generate-proposal", "--topic", "test topic", "--language", "en"])
    assert args4.command == "generate-proposal"
    args5 = parser.parse_args(["generate-presentation", "--topic", "test topic", "--deck-type", "lab_update"])
    assert args5.deck_type == "lab_update"
    args6 = parser.parse_args(["generate-paper-draft", "--topic", "test topic", "--paper-type", "conference"])
    assert args6.paper_type == "conference"
    args7 = parser.parse_args(["generate-literature-review", "--topic", "test topic", "--language", "en"])
    assert args7.command == "generate-literature-review"
    args8 = parser.parse_args(["analyze-capabilities"])
    assert args8.command == "analyze-capabilities"
    args9 = parser.parse_args(["answer-research-question", "--question", "what methods are common?", "--language", "en"])
    assert args9.command == "answer-research-question"
    args10 = parser.parse_args(["export-docx", "--artifact", "paper"])
    assert args10.command == "export-docx"
    args11 = parser.parse_args(["export-pptx"])
    assert args11.command == "export-pptx"
test("research_bridge CLI parser", test_research_bridge_parser)


def test_research_bridge_refresh():
    import shutil
    import uuid
    from pathlib import Path
    from tools.research_bridge import main

    root = Path("output") / f"tmp-bridge-test-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        (root / "knowledge-base").mkdir()
        (root / "drafts").mkdir()
        (root / "output").mkdir()
        (root / "papers").mkdir()

        exit_code = main(["refresh", "--project-root", str(root)])
        assert exit_code == 0
    finally:
        shutil.rmtree(root)
test("research_bridge refresh", test_research_bridge_refresh)


def test_research_bridge_capability_audit():
    import shutil
    import uuid
    from pathlib import Path
    from tools.research_bridge import main

    root = Path("output") / f"tmp-capability-audit-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        (root / "knowledge-base").mkdir()
        (root / "drafts").mkdir()
        (root / "output").mkdir()
        (root / "papers").mkdir()
        (root / "knowledge-base" / "paper_index.json").write_text("[]", encoding="utf-8")

        exit_code = main(["analyze-capabilities", "--project-root", str(root)])
        assert exit_code == 0
        assert (root / "drafts" / "research-capability-audit.md").exists()
    finally:
        shutil.rmtree(root)
test("research_bridge capability audit", test_research_bridge_capability_audit)


def test_research_bridge_literature_review():
    import shutil
    import uuid
    import json
    from pathlib import Path
    from tools.research_bridge import main

    root = Path("output") / f"tmp-literature-review-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        (root / "knowledge-base").mkdir()
        (root / "drafts").mkdir()
        (root / "output").mkdir()
        (root / "papers").mkdir()
        (root / "knowledge-base" / "paper_index.json").write_text("[]", encoding="utf-8")

        exit_code = main(["generate-literature-review", "--project-root", str(root), "--topic", "test topic", "--language", "en"])
        assert exit_code == 0
        assert (root / "drafts" / "literature-review.md").exists()
        payload = json.loads((root / "output" / "literature-review.json").read_text(encoding="utf-8"))
        assert len(payload["sections"]) >= 3
    finally:
        shutil.rmtree(root)
test("research_bridge literature review", test_research_bridge_literature_review)


def test_research_bridge_research_qa():
    import shutil
    import uuid
    import json
    from pathlib import Path
    from tools.research_bridge import main

    root = Path("output") / f"tmp-research-qa-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        (root / "knowledge-base").mkdir()
        (root / "drafts").mkdir()
        (root / "output").mkdir()
        (root / "papers").mkdir()
        (root / "output" / "paper-content").mkdir(parents=True)
        (root / "output" / "paper-content" / "demo.json").write_text(
            json.dumps(
                {
                    "content": (
                        "Graph neural networks are widely used for structured prediction.\n\n"
                        "This study evaluates graph neural network encoders on public benchmarks.\n\n"
                        "The experiments report accuracy, latency, and robustness on OGB datasets."
                    )
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        (root / "knowledge-base" / "paper_index.json").write_text(
            json.dumps(
                [
                    {
                        "record_id": "paper:test",
                        "title": "Graph neural network evaluation",
                        "abstract": "We benchmark graph neural network methods on public datasets.",
                        "content_json_path": "output/paper-content/demo.json",
                        "content_crawled": True,
                        "year": 2025,
                        "source": "local",
                    }
                ],
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        exit_code = main(["answer-research-question", "--project-root", str(root), "--question", "what are common methods?", "--language", "en"])
        assert exit_code == 0
        assert (root / "drafts" / "research-answer.md").exists()
        payload = json.loads((root / "output" / "research-answer.json").read_text(encoding="utf-8"))
        assert payload["answer_blocks"]
        assert payload["evidence"][0]["top_chunks"]
    finally:
        shutil.rmtree(root)
test("research_bridge research qa", test_research_bridge_research_qa)


def test_research_bridge_export_docx():
    import shutil
    import uuid
    from pathlib import Path
    from tools.research_bridge import main

    root = Path("output") / f"tmp-export-docx-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        (root / "knowledge-base").mkdir()
        (root / "drafts").mkdir()
        (root / "output").mkdir()
        (root / "papers").mkdir()
        (root / "drafts" / "paper-draft.md").write_text("# Test Draft\n\nA short paragraph.\n", encoding="utf-8")

        exit_code = main(["export-docx", "--project-root", str(root), "--artifact", "paper"])
        assert exit_code == 0
        assert list((root / "output" / "exports").glob("*.docx"))
    finally:
        shutil.rmtree(root)
test("research_bridge export docx", test_research_bridge_export_docx)


def test_research_bridge_export_docx_lazy_imports_without_numpy():
    import argparse
    import importlib
    import shutil
    import sys
    import types
    import uuid
    from pathlib import Path

    root = Path("output") / f"tmp-export-docx-lazy-{uuid.uuid4().hex}"
    root.mkdir(parents=True)

    old_bridge = sys.modules.pop("tools.research_bridge", None)
    old_data = sys.modules.get("tools.data_analyzer")
    old_fig = sys.modules.get("tools.figure_generator")
    trap_hits: list[tuple[str, str]] = []

    def _trap_module(name: str):
        module = types.ModuleType(name)

        def __getattr__(attr: str):
            trap_hits.append((name, attr))
            raise AssertionError(f"{name} should not be touched during DOCX export")

        module.__getattr__ = __getattr__  # type: ignore[attr-defined]
        return module

    sys.modules["tools.data_analyzer"] = _trap_module("tools.data_analyzer")
    sys.modules["tools.figure_generator"] = _trap_module("tools.figure_generator")

    try:
        bridge = importlib.import_module("tools.research_bridge")

        (root / "knowledge-base").mkdir()
        (root / "drafts").mkdir()
        (root / "output").mkdir()
        (root / "papers").mkdir()
        (root / "drafts" / "paper-draft.md").write_text("# Test Draft\n\nA short paragraph.\n", encoding="utf-8")

        original_export = bridge.export_markdown_to_docx
        original_sync = bridge.sync_project_state
        original_dashboard = bridge.build_dashboard
        try:
            bridge.export_markdown_to_docx = lambda **kwargs: {"markdown_path": kwargs.get("source") or "drafts/paper-draft.md"}
            bridge.sync_project_state = lambda project_root: {"stage": "export"}
            bridge.build_dashboard = lambda project_root: Path(project_root) / "output" / "paper-workbench.html"

            result = bridge._export_docx(
                argparse.Namespace(
                    project_root=str(root),
                    artifact="paper",
                    source=None,
                    output=None,
                    docx_style="default",
                )
            )

            assert result["artifact"]["markdown_path"]
            assert trap_hits == []
        finally:
            bridge.export_markdown_to_docx = original_export
            bridge.sync_project_state = original_sync
            bridge.build_dashboard = original_dashboard
    finally:
        sys.modules.pop("tools.research_bridge", None)
        if old_bridge is not None:
            sys.modules["tools.research_bridge"] = old_bridge
        if old_data is None:
            sys.modules.pop("tools.data_analyzer", None)
        else:
            sys.modules["tools.data_analyzer"] = old_data
        if old_fig is None:
            sys.modules.pop("tools.figure_generator", None)
        else:
            sys.modules["tools.figure_generator"] = old_fig
        shutil.rmtree(root, ignore_errors=True)
test("research_bridge export docx avoids numpy-heavy imports", test_research_bridge_export_docx_lazy_imports_without_numpy)


def test_research_bridge_export_pptx():
    import shutil
    import uuid
    import json
    from pathlib import Path
    from tools.research_bridge import main

    root = Path("output") / f"tmp-export-pptx-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        (root / "knowledge-base").mkdir()
        (root / "drafts").mkdir()
        (root / "output").mkdir()
        (root / "papers").mkdir()
        (root / "output" / "research-presentation.json").write_text(
            json.dumps(
                {
                    "title": "Test Presentation",
                    "slides": [
                        {"title": "Slide 1", "bullets": ["One", "Two"], "notes": "Speaker note"},
                        {"title": "Slide 2", "bullets": ["Three"], "notes": "Another note"},
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        exit_code = main(["export-pptx", "--project-root", str(root)])
        assert exit_code == 0
        assert list((root / "output" / "exports").glob("*.pptx"))
    finally:
        shutil.rmtree(root)
test("research_bridge export pptx", test_research_bridge_export_pptx)

def test_writing_enhancer_section_evidence_packet_surfaces_figure_slots():
    from tools.writing_enhancer import _build_section_evidence_packet

    packet = _build_section_evidence_packet(
        topic="robot navigation",
        language="en",
        section_title="4. Experiment Results",
        section_points=["Result analysis"],
        project_context={
            "figure_plan": [
                {
                    "caption": "Figure 4-1 Runtime comparison",
                    "goal": "Compare planning latency",
                    "figure_type": "bar",
                    "evidence": "results/latency.csv",
                    "existing_asset": "output/figures/latency.png",
                }
            ]
        },
        project_brief="project brief line",
        reference_brief="reference cue line",
    )

    assert "Figure/table anchors" in packet
    assert "Runtime comparison" in packet
    assert "latency.png" in packet
    assert "project brief line" in packet
    assert "reference cue line" in packet
test("writing enhancer evidence packet includes figure slots", test_writing_enhancer_section_evidence_packet_surfaces_figure_slots)


def test_writing_enhancer_expand_short_sections_reviews_before_rewrite():
    import tools.writing_enhancer as enhancer

    prompts: list[str] = []
    original_call = enhancer._call_model

    try:
        def fake_call_model(config, prompt, language):
            prompts.append(prompt)
            if "Respond with JSON only" in prompt:
                return json.dumps(
                    {
                        "score": 62,
                        "issues": ["needs stronger evidence anchoring"],
                        "missing_evidence": ["connect the claim to Figure 4-1"],
                        "rewrite_plan": ["add one paragraph on latency trade-offs"],
                        "preserve": ["keep the subsection structure"],
                        "figure_actions": ["tie the main claim to the runtime comparison slot"],
                    }
                )
            return (
                "### Result analysis\n"
                "The section now anchors the runtime discussion to Figure 4-1 and explains the latency trade-off.\n\n"
                "A second paragraph adds implementation-specific detail about queueing, planner updates, and constraints."
            )

        enhancer._call_model = fake_call_model
        sections = enhancer._expand_short_sections(
            llm_config={"provider": "test", "model": "test", "base_url": "http://example.invalid", "api_key": "x"},
            topic="robot navigation",
            language="en",
            references=[],
            project_context={
                "figure_plan": [
                    {
                        "caption": "Figure 4-1 Runtime comparison",
                        "goal": "Compare planning latency",
                        "figure_type": "bar",
                        "evidence": "results/latency.csv",
                        "existing_asset": "output/figures/latency.png",
                    }
                ]
            },
            sections=[
                {
                    "title": "4. Experiment Results",
                    "content": ["### Result analysis", "Short baseline discussion."],
                }
            ],
            target_words=1200,
            remaining_gap=500,
            target_max_words=1400,
            blueprint=[{"share": 1.0}],
        )

        assert len(prompts) == 2
        assert "Respond with JSON only" in prompts[0]
        assert "Review feedback:" in prompts[1]
        assert "Figure/table anchors" in prompts[1]
        assert "Runtime comparison" in prompts[1]
        assert sections[0]["content"][1].startswith("The section now anchors")
    finally:
        enhancer._call_model = original_call
test("writing enhancer expansion reviews before rewrite", test_writing_enhancer_expand_short_sections_reviews_before_rewrite)


def test_writing_enhancer_quality_meta_marks_review_pipeline():
    import shutil
    import uuid
    from pathlib import Path
    import tools.writing_enhancer as enhancer

    root = Path("output") / f"tmp-writing-enhancer-quality-meta-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        markdown_path = root / "draft.md"
        outline_path = root / "outline.md"
        plan_path = root / "plan.md"
        json_path = root / "draft.json"
        base_result = {
            "markdown_path": str(markdown_path),
            "outline_path": str(outline_path),
            "plan_path": str(plan_path),
            "json_path": str(json_path),
            "artifact": {"title": "quality meta test", "references": []},
        }

        original_load = enhancer._load_llm_config
        original_draft = enhancer._draft_sections_with_llm
        original_expand = enhancer._expand_short_sections

        try:
            enhancer._load_llm_config = lambda *args, **kwargs: [
                {"provider": "test", "model": "test-model", "base_url": "http://example.invalid", "api_key": "x"}
            ]
            enhancer._draft_sections_with_llm = lambda **kwargs: [
                {"title": "1. Introduction", "content": ["### Background", "a" * 1800]},
                {"title": "2. Results", "content": ["### Result analysis", "b" * 1800]},
            ]
            enhancer._expand_short_sections = lambda **kwargs: kwargs["sections"]

            result = enhancer.enhance_generated_paper_package(
                base_result,
                project_root=root,
                topic="quality meta test",
                language="en",
                paper_type="general",
                project_context={"figure_plan": [{"caption": "Figure 2-1", "goal": "show result"}]},
                target_words=3000,
            )

            quality_meta = result["artifact"]["quality_meta"]
            assert quality_meta["section_evidence_packets"] is True
            assert quality_meta["critique_revision_enabled"] is True
            assert quality_meta["quality_review_threshold"] >= 80
            assert quality_meta["figure_plan_anchored"] is True
        finally:
            enhancer._load_llm_config = original_load
            enhancer._draft_sections_with_llm = original_draft
            enhancer._expand_short_sections = original_expand
    finally:
        shutil.rmtree(root)
test("writing enhancer quality meta exposes review pipeline", test_writing_enhancer_quality_meta_marks_review_pipeline)

# ============================================================
print("\n=== 8. 写作增强回归 (本地) ===")
# ============================================================

def test_writing_enhancer_project_brief_localization():
    from tools.writing_enhancer import _build_project_brief

    zh = _build_project_brief(None, "zh")
    en = _build_project_brief(None, "en")

    assert "No direct project evidence was supplied." not in zh
    assert "项目实现证据" in zh
    assert en == "No direct project evidence was supplied."
test("writing enhancer zh project brief localization", test_writing_enhancer_project_brief_localization)


def test_writing_enhancer_zh_spacing_and_fallback():
    from tools.writing_enhancer import _build_fallback_sections, _strip_ai_tone

    cleaned = _strip_ai_tone(
        "结合当前项目证据，可用线索包括： No direct project evidence was supplied. 写作时应参考 output / exports / paper-draft.docx 。",
        "zh",
    )
    assert "No direct project evidence was supplied." in cleaned
    assert "output / exports / paper-draft.docx" in cleaned

    sections = _build_fallback_sections(
        topic="基于ROS的导航系统",
        language="zh",
        references=[],
        project_context=None,
        blueprint=[{"title": "1. 绪论", "points": ["研究背景与意义"]}],
    )
    joined = "\n".join("\n".join(section["content"]) for section in sections)
    assert "No direct project evidence was supplied." not in joined
    assert "当前尚未提供可直接引用的项目实现证据" in joined
test("writing enhancer zh spacing and fallback", test_writing_enhancer_zh_spacing_and_fallback)


def test_writing_enhancer_section_aware_briefs():
    from tools.writing_enhancer import _build_project_brief, _build_reference_brief

    project_context = {
        "project_name": "demo-nav",
        "project_summary": "A ROS2 navigation stack with mapping and planning modules.",
        "stack": ["Python", "ROS2"],
        "candidate_source_files": ["src/planner.py", "src/controller.py", "src/eval.py"],
        "candidate_config_files": ["config/nav2_params.yaml", "config/slam.yaml"],
        "candidate_result_files": ["output/eval_metrics.csv", "output/trajectory_report.md"],
        "method_clues": ["function `plan_path` in `planner.py`", "class `ControllerNode` in `controller.py`"],
        "result_clues": ["Potential result evidence in `eval_metrics.csv`", "Potential result evidence in `trajectory_report.md`"],
    }

    method_brief = _build_project_brief(
        project_context,
        "en",
        topic="robot navigation",
        section_title="3. Method and System Design",
        section_points=["Architecture", "Interfaces"],
    )
    experiment_brief = _build_project_brief(
        project_context,
        "en",
        topic="robot navigation",
        section_title="5. Experiments and Analysis",
        section_points=["Metrics", "Results"],
    )
    assert "planner.py" in method_brief
    assert "Result clues:" in experiment_brief
    assert "eval_metrics.csv" in experiment_brief

    references = [
        {
            "title": "Visual SLAM system design for indoor robots",
            "abstract": "Architecture and mapping pipeline for robot navigation.",
            "year": 2022,
        },
        {
            "title": "Benchmarking robot navigation metrics",
            "abstract": "Evaluation metrics, experiments, and benchmark analysis for mobile robot navigation.",
            "year": 2024,
        },
        {
            "title": "Unrelated crop disease survey",
            "abstract": "Agricultural disease datasets and models.",
            "year": 2025,
        },
    ]
    experiment_refs = _build_reference_brief(
        references,
        topic="robot navigation",
        section_title="5. Experiments and Analysis",
        section_points=["Metrics", "Results"],
        language="en",
    )
    assert "Benchmarking robot navigation metrics" in experiment_refs.splitlines()[0]
test("writing enhancer section-aware briefs", test_writing_enhancer_section_aware_briefs)


def test_writing_enhancer_prompt_guardrails_and_structure():
    from tools.writing_enhancer import _build_section_guardrails, _build_section_prompt, _sanitize_section_output

    notes = _build_section_guardrails(
        language="en",
        section_title="5. Experiments and Analysis",
        section_points=["Setup", "Metrics", "Results"],
        has_project_context=True,
        has_result_evidence=False,
        has_reference_support=False,
    )
    assert "figure/table slots" in notes

    prompt = _build_section_prompt(
        topic="robot navigation",
        language="en",
        paper_type="thesis",
        section_title="5. Experiments and Analysis",
        section_points=["Setup", "Metrics", "Results"],
        section_target=900,
        reference_brief="No strong indexed references are currently available.",
        project_brief="Direct result evidence is still missing, so this section should stay at the level of setup.",
        section_notes=notes,
    )
    assert "### Setup" in prompt
    assert "### Metrics" in prompt
    assert "Section guardrails" in prompt

    cleaned = _sanitize_section_output(
        "The evaluation section discusses setup and metrics.\n\n- latency\n- accuracy\n- robustness",
        "5. Experiments and Analysis",
        "en",
        section_points=["Setup", "Metrics", "Results"],
    )
    assert "### Setup" in cleaned
    assert "- latency" not in cleaned
    assert "latency; accuracy; robustness." in cleaned
test("writing enhancer prompt guardrails", test_writing_enhancer_prompt_guardrails_and_structure)


def test_writing_enhancer_zh_rhetorical_prefix_cleanup():
    from tools.writing_enhancer import _strip_ai_tone

    cleaned = _strip_ai_tone(
        "结果表明，SLAM Toolbox 在地图精度上保持较高水平。综合来看，系统在典型室内场景下运行稳定。需要指出的是，参数整定仍依赖经验。",
        "zh",
    )
    assert "结果表明" not in cleaned
    assert "综合来看" not in cleaned
    assert "需要指出的是" not in cleaned
    assert "SLAM Toolbox 在地图精度上保持较高水平" in cleaned
    assert "参数整定仍依赖经验" in cleaned
test("writing enhancer zh rhetorical prefix cleanup", test_writing_enhancer_zh_rhetorical_prefix_cleanup)


def test_writing_enhancer_root_base_url_uses_v1_chat_path():
    import tools.writing_enhancer as enhancer

    class DummyResponse:
        def __init__(self, url, payload):
            self.url = url
            self._payload = payload
            self.text = ""
            self.status_code = 200
            self.headers = {"content-type": "application/json"}

        def raise_for_status(self):
            return None

        def json(self):
            return self._payload

    calls: list[str] = []
    original_post = enhancer.requests.post
    try:
        def fake_post(url, **kwargs):
            calls.append(url)
            return DummyResponse(url, {"choices": [{"message": {"content": "ok"}}]})

        enhancer.requests.post = fake_post
        result = enhancer._call_model(
            {
                "base_url": "https://www.xiangluapi.com",
                "provider": "llm",
                "model": "gemini-3.1-pro-high",
                "api_key": "dummy",
            },
            "hello",
            "zh",
        )
        assert result == "ok"
        assert calls[-1].endswith("/v1/chat/completions")
    finally:
        enhancer.requests.post = original_post
test("writing enhancer root base url uses v1 chat path", test_writing_enhancer_root_base_url_uses_v1_chat_path)


def test_writing_enhancer_plain_text_response_fallback():
    import tools.writing_enhancer as enhancer

    class DummyResponse:
        def __init__(self, url, text):
            self.url = url
            self.text = text
            self.status_code = 200
            self.headers = {"content-type": "text/plain; charset=utf-8"}

        def raise_for_status(self):
            return None

        def json(self):
            raise ValueError("no json body")

    original_post = enhancer.requests.post
    try:
        enhancer.requests.post = lambda url, **kwargs: DummyResponse(url, "plain fallback answer")
        result = enhancer._call_model(
            {
                "base_url": "https://www.xiangluapi.com",
                "provider": "llm",
                "model": "gemini-3.1-pro-high",
                "api_key": "dummy",
            },
            "hello",
            "zh",
        )
        assert result == "plain fallback answer"
    finally:
        enhancer.requests.post = original_post
test("writing enhancer plain text response fallback", test_writing_enhancer_plain_text_response_fallback)


def test_text_safety_sanitize_utf8_text_handles_surrogates():
    from tools.text_safety import contains_surrogates, sanitize_utf8_text

    repaired_pair = sanitize_utf8_text("A\ud835\udc00B")
    repaired_orphan = sanitize_utf8_text("A\ud835B")

    assert repaired_pair == f"A{chr(0x1D400)}B"
    assert not contains_surrogates(repaired_pair)
    assert not contains_surrogates(repaired_orphan)
test("text safety sanitizes surrogate text", test_text_safety_sanitize_utf8_text_handles_surrogates)


def test_writing_enhancer_call_model_sanitizes_surrogate_payloads():
    import tools.writing_enhancer as enhancer
    from tools.text_safety import contains_surrogates

    captured_payloads: list[dict] = []

    class DummyResponse:
        def __init__(self, url, payload):
            self.url = url
            self._payload = payload
            self.text = json.dumps(payload, ensure_ascii=False)
            self.status_code = 200
            self.headers = {"content-type": "application/json; charset=utf-8"}

        def raise_for_status(self):
            return None

        def json(self):
            return self._payload

    original_post = enhancer.requests.post
    try:
        def fake_post(url, **kwargs):
            captured_payloads.append(kwargs["json"])
            return DummyResponse(url, {"choices": [{"message": {"content": "ok\ud835"}}]})

        enhancer.requests.post = fake_post
        result = enhancer._call_model(
            {
                "base_url": "https://www.xiangluapi.com",
                "provider": "llm",
                "model": "gemini-3.1-pro-high",
                "api_key": "dummy",
            },
            "hello\ud835world",
            "zh",
        )

        assert captured_payloads
        for message in captured_payloads[0]["messages"]:
            assert not contains_surrogates(message["content"])
        assert not contains_surrogates(result)
        assert result.startswith("ok")
    finally:
        enhancer.requests.post = original_post
test("writing enhancer sanitizes surrogate payloads", test_writing_enhancer_call_model_sanitizes_surrogate_payloads)


def test_paper_writer_prompt_assets_include_anti_template_rules():
    from tools.paper_writer import _render_clean_zh_revision_prompts_markdown, _render_writing_assets_markdown

    prompts = _render_clean_zh_revision_prompts_markdown("示例题目")
    assets = _render_writing_assets_markdown()

    assert "初稿去模板味改写" in prompts
    assert "结果表明、由此可见、综合来看" in prompts
    assert "句式雷同扫描" in prompts
    assert "项目证据优先" in assets
    assert "高频句壳" in assets
test("paper writer anti-template prompt assets", test_paper_writer_prompt_assets_include_anti_template_rules)


def test_writing_enhancer_cross_section_dedup_and_fallback_focus():
    from tools.writing_enhancer import _build_fallback_sections, _dedupe_sections

    sections = [
        {
            "title": "2. Technical Background",
            "content": [
                "### Core concepts",
                "The subsection on Core concepts should explain its role in the overall study, clarify the design constraints, and connect technical choices to the larger research objective.",
            ],
        },
        {
            "title": "4. Implementation",
            "content": [
                "### Core modules",
                "The subsection on Core modules should explain its role in the overall study, clarify the design constraints, and connect technical choices to the larger research objective.",
                "The subsection on Core modules should stay close to code structure, runtime flow, configuration, and module coordination rather than repeating high-level terminology.",
            ],
        },
    ]
    deduped = _dedupe_sections(sections, "en")
    joined = "\n".join(deduped[1]["content"])
    assert "should explain its role in the overall study" not in joined
    assert "code structure, runtime flow, configuration, and module coordination" in joined

    fallback = _build_fallback_sections(
        topic="robot navigation",
        language="en",
        references=[],
        project_context=None,
        blueprint=[{"title": "4. Implementation", "points": ["Core modules"]}],
    )
    fallback_joined = "\n".join(fallback[0]["content"])
    assert "code structure, runtime flow, configuration, and module coordination" in fallback_joined
    assert "should explain its role in the overall study" not in fallback_joined
test("writing enhancer cross-section dedup", test_writing_enhancer_cross_section_dedup_and_fallback_focus)


def test_writing_enhancer_experiment_placeholder_standardization():
    from tools.writing_enhancer import _build_fallback_sections, _build_project_brief, _build_section_guardrails

    project_context = {
        "project_name": "demo-nav",
        "project_summary": "A ROS2 navigation stack with mapping and planning modules.",
        "stack": ["Python", "ROS2"],
        "candidate_source_files": ["src/planner.py"],
        "candidate_config_files": ["config/nav2_params.yaml"],
        "method_clues": ["function `plan_path` in `planner.py`"],
        "result_clues": [],
        "candidate_result_files": [],
    }

    brief = _build_project_brief(
        project_context,
        "en",
        topic="robot navigation",
        section_title="5. Experiments and Analysis",
        section_points=["Setup", "Results"],
    )
    assert "Direct result evidence is currently unavailable" in brief

    notes = _build_section_guardrails(
        language="en",
        section_title="5. Experiments and Analysis",
        section_points=["Setup", "Results"],
        has_project_context=True,
        has_result_evidence=False,
        has_reference_support=True,
    )
    assert "Direct result evidence is currently unavailable" in notes

    fallback = _build_fallback_sections(
        topic="robot navigation",
        language="en",
        references=[],
        project_context=project_context,
        blueprint=[{"title": "5. Experiments and Analysis", "points": ["Setup"]}],
    )
    joined = "\n".join(fallback[0]["content"])
    assert "Direct result evidence is currently unavailable" in joined
    assert "evaluation environment, task setting, metrics, baselines, and figure/table slots" in joined
test("writing enhancer experiment placeholder standardization", test_writing_enhancer_experiment_placeholder_standardization)


def test_writing_enhancer_low_signal_sentence_pruning():
    from tools.writing_enhancer import _sanitize_section_output

    cleaned = _sanitize_section_output(
        "This section discusses 4. Implementation for the topic robot navigation, keeping the argument focused on motivation, implementation logic, and evidence-backed analysis. "
        "The planner module reads nav2_params.yaml and publishes commands to the controller at runtime.",
        "4. Implementation",
        "en",
        section_points=["Core modules"],
    )
    assert "This section discusses" not in cleaned
    assert "planner module reads nav2_params.yaml" in cleaned

    experiment_cleaned = _sanitize_section_output(
        "Direct result evidence is currently unavailable, so this section should keep only evaluation setup, metrics, baselines, and figure/table slots. "
        "Table 1 will be reserved for the main benchmark summary after logs are collected.",
        "5. Experiments and Analysis",
        "en",
        section_points=["Setup", "Results"],
    )
    assert "Direct result evidence is currently unavailable" in experiment_cleaned
    assert "Table 1 will be reserved for the main benchmark summary" in experiment_cleaned
test("writing enhancer low-signal pruning", test_writing_enhancer_low_signal_sentence_pruning)


def test_writing_enhancer_preserves_existing_base_sections_without_llm():
    import json
    import shutil
    import uuid
    from pathlib import Path
    from tools.writing_enhancer import enhance_generated_paper_package

    root = Path("output") / f"tmp-writing-enhancer-preserve-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        markdown_path = root / "draft.md"
        outline_path = root / "outline.md"
        plan_path = root / "plan.md"
        json_path = root / "draft.json"
        base_result = {
            "markdown_path": str(markdown_path),
            "outline_path": str(outline_path),
            "plan_path": str(plan_path),
            "json_path": str(json_path),
            "artifact": {
                "title": "示例题目",
                "sections": [
                    {
                        "title": "1. 引言",
                        "content": [
                            "### 研究背景",
                            "该系统围绕 ROS2 导航链路构建，已经包含建图、定位与控制模块的工程实现。",
                        ],
                    }
                ],
                "references": [],
            },
        }

        result = enhance_generated_paper_package(
            base_result,
            project_root=root,
            topic="示例题目",
            language="zh",
            paper_type="general",
            project_context=None,
            target_words=6000,
        )
        content = markdown_path.read_text(encoding="utf-8")
        payload = json.loads(json_path.read_text(encoding="utf-8"))
        assert "已经包含建图、定位与控制模块的工程实现" in content
        assert payload["sections"][0]["content"][1].startswith("该系统围绕 ROS2 导航链路构建")
        assert result["artifact"]["quality_meta"]["base_sections_preserved"] is True
    finally:
        shutil.rmtree(root)
test("writing enhancer preserves base sections", test_writing_enhancer_preserves_existing_base_sections_without_llm)


def test_writing_enhancer_high_target_keeps_expanding_until_near_goal():
    import shutil
    import uuid
    from pathlib import Path
    import tools.writing_enhancer as enhancer

    root = Path("output") / f"tmp-writing-enhancer-target-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        markdown_path = root / "draft.md"
        outline_path = root / "outline.md"
        plan_path = root / "plan.md"
        json_path = root / "draft.json"
        base_result = {
            "markdown_path": str(markdown_path),
            "outline_path": str(outline_path),
            "plan_path": str(plan_path),
            "json_path": str(json_path),
            "artifact": {
                "title": "高字数目标测试",
                "references": [],
            },
        }

        original_load = enhancer._load_llm_config
        original_draft = enhancer._draft_sections_with_llm
        original_expand = enhancer._expand_short_sections
        expand_calls = {"count": 0}

        try:
            enhancer._load_llm_config = lambda *args, **kwargs: [
                {"provider": "test", "model": "test", "base_url": "http://example.invalid", "api_key": "x"}
            ]

            def fake_draft_sections_with_llm(**kwargs):
                return [
                    {"title": "1. 绪论", "content": ["### 背景", "甲" * 2500]},
                    {"title": "2. 技术基础", "content": ["### 理论", "乙" * 2600]},
                    {"title": "3. 系统设计", "content": ["### 架构", "丙" * 2800]},
                    {"title": "4. 关键实现", "content": ["### 模块", "丁" * 3000]},
                    {"title": "5. 实验分析", "content": ["### 结果", "戊" * 3200]},
                    {"title": "6. 结论", "content": ["### 总结", "己" * 1600]},
                ]

            def fake_expand_short_sections(**kwargs):
                expand_calls["count"] += 1
                sections = []
                for section in kwargs["sections"]:
                    sections.append({"title": section["title"], "content": list(section["content"])})
                sections[0]["content"].append(f"第{expand_calls['count']}轮扩写" + ("扩" * 1500))
                return sections

            enhancer._draft_sections_with_llm = fake_draft_sections_with_llm
            enhancer._expand_short_sections = fake_expand_short_sections

            result = enhancer.enhance_generated_paper_package(
                base_result,
                project_root=root,
                topic="高字数目标测试",
                language="zh",
                paper_type="general",
                project_context=None,
                target_words=20000,
            )

            assert result["artifact"]["actual_words"] >= 19000
            assert result["artifact"]["quality_meta"]["expansion_rounds"] >= 2
            assert expand_calls["count"] >= 2
        finally:
            enhancer._load_llm_config = original_load
            enhancer._draft_sections_with_llm = original_draft
            enhancer._expand_short_sections = original_expand
    finally:
        shutil.rmtree(root)
test("writing enhancer high target expansion loop", test_writing_enhancer_high_target_keeps_expanding_until_near_goal)


def test_writing_enhancer_hits_requested_floor_without_large_overshoot():
    import shutil
    import uuid
    from pathlib import Path
    import tools.writing_enhancer as enhancer

    root = Path("output") / f"tmp-writing-enhancer-target-window-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        markdown_path = root / "draft.md"
        outline_path = root / "outline.md"
        plan_path = root / "plan.md"
        json_path = root / "draft.json"
        base_result = {
            "markdown_path": str(markdown_path),
            "outline_path": str(outline_path),
            "plan_path": str(plan_path),
            "json_path": str(json_path),
            "artifact": {
                "title": "target window test",
                "references": [],
            },
        }

        original_load = enhancer._load_llm_config
        original_draft = enhancer._draft_sections_with_llm
        original_expand = enhancer._expand_short_sections

        try:
            enhancer._load_llm_config = lambda *args, **kwargs: [
                {"provider": "test", "model": "test", "base_url": "http://example.invalid", "api_key": "x"}
            ]

            def fake_draft_sections_with_llm(**kwargs):
                return [
                    {"title": "1. Introduction", "content": ["### Background", "a" * 2500]},
                    {"title": "2. Foundation", "content": ["### Theory", "b" * 2600]},
                    {"title": "3. Design", "content": ["### Architecture", "c" * 2800]},
                    {"title": "4. Implementation", "content": ["### Modules", "d" * 3000]},
                    {"title": "5. Experiments", "content": ["### Results", "e" * 3200]},
                    {"title": "6. Conclusion", "content": ["### Summary", "f" * 1600]},
                ]

            enhancer._draft_sections_with_llm = fake_draft_sections_with_llm

            for target_words in (12000, 20000, 35000):
                expand_calls = {"count": 0}

                def fake_expand_short_sections(**kwargs):
                    expand_calls["count"] += 1
                    remaining_gap = int(kwargs.get("remaining_gap") or 0)
                    threshold = enhancer._final_fill_threshold(kwargs["target_words"], "zh")
                    if remaining_gap <= threshold:
                        addition = remaining_gap + 1200
                    else:
                        addition = max(1800, int(remaining_gap * 0.55))
                    sections = []
                    for section in kwargs["sections"]:
                        sections.append({"title": section["title"], "content": list(section["content"])})
                    sections[0]["content"].append(
                        f"round-{target_words}-{expand_calls['count']}-" + ("扩" * addition)
                    )
                    return sections

                enhancer._expand_short_sections = fake_expand_short_sections

                result = enhancer.enhance_generated_paper_package(
                    base_result,
                    project_root=root,
                    topic="target window test",
                    language="zh",
                    paper_type="general",
                    project_context=None,
                    target_words=target_words,
                )

                actual_words = result["artifact"]["actual_words"]
                assert actual_words >= target_words
                assert actual_words <= enhancer._target_max_words(target_words, "zh")
                assert result["artifact"]["quality_meta"]["expansion_rounds"] == expand_calls["count"]
        finally:
            enhancer._load_llm_config = original_load
            enhancer._draft_sections_with_llm = original_draft
            enhancer._expand_short_sections = original_expand
    finally:
        shutil.rmtree(root)
test("writing enhancer target floor and cap", test_writing_enhancer_hits_requested_floor_without_large_overshoot)

# ============================================================
print("\n=== 9. API 联通性 (网络) ===")
# ============================================================

def test_paper_writer_zh_polish_drops_summary_blocks_and_shells():
    from tools.paper_writer import _polish_chinese_sections

    sections = [
        {
            "title": "1. 绪论",
            "content": [
                "### 1.6 本章小结",
                "本章围绕课题背景、研究意义和技术路线展开论述，后续工作主要集中在补充实验图表与文献引用。",
                "### 1.7 研究现状、评价维度与本文切入点",
                "从真实机器人参数文件可以看出，系统配置了 AMCL、控制器服务器与代价地图。",
            ],
        }
    ]

    polished = _polish_chinese_sections(sections)
    content = "\n".join(polished[0]["content"])
    assert "本章小结" not in content
    assert "后续工作主要集中在补充实验图表" not in content
    assert "本文切入点" not in content
    assert "课题切入点" in content
    assert "从真实机器人参数文件可以看出" not in content
    assert "真实机器人参数文件中" in content
test("paper writer zh polish removes summary shells", test_paper_writer_zh_polish_drops_summary_blocks_and_shells)


def test_paper_writer_zh_polish_rewrites_project_doc_tone():
    from tools.paper_writer import _rewrite_zh_paragraph_surface

    rewritten = _rewrite_zh_paragraph_surface(
        "该项目是一个面向移动机器人的 Python、C/C++ 工程，包含建图、导航、雷达驱动与整车启动相关模块。",
        section_title="1. 绪论",
        heading="1.2 课题意义",
    )
    assert rewritten.startswith("项目代码以 Python、C/C++ 为主")
    assert "该项目是一个面向移动机器人的" not in rewritten
test("paper writer zh polish rewrites project-doc tone", test_paper_writer_zh_polish_rewrites_project_doc_tone)


def test_paper_writer_zh_polish_rewrites_remaining_meta_phrases():
    from tools.paper_writer import _polish_chinese_sections

    sections = [
        {
            "title": "4. 系统实现",
            "content": [
                "### 4.5 调试方法与实现小结",
                "对于工程型课题而言，评价重点不能局限于单一算法指标，而应同时覆盖系统功能完整性、运行稳定性、模块可维护性、实验可重复性与应用扩展能力等多个维度。只有在这些维度上形成较为均衡的表现，系统方案才具备较强的工程价值与推广意义。因此，本文后续的系统设计与实验章节将重点围绕这些评价维度展开，从而体现本课题在工程实现与综合设计上的主要贡献。",
                "通过把参数配置与控制效果之间建立明确对应关系，可以让论文从“功能说明文档”转变为“设计决策说明文档”，这对于本科毕业论文尤其重要。",
                "受现有项目材料限制，本章先根据已有实现与运行链路展开分析，并将实验指标、图表位置和结果解释框架固定下来。在补充完整实验记录后，这种组织方式能够较为自然地衔接定量结果与文字分析。",
            ],
        }
    ]

    polished = _polish_chinese_sections(sections)
    content = "\n".join(polished[0]["content"])
    assert "调试方法与实现小结" not in content
    assert "调试方法与实现分析" in content
    assert "本文后续的系统设计与实验章节" not in content
    assert "后续系统设计与实验分析需要围绕这些评价维度展开" in content
    assert "功能说明文档" not in content
    assert "设计决策说明文档" not in content
    assert "参数整定的工程依据" in content
    assert "受现有项目材料限制" not in content
    assert "评价指标、实验场景和结果解释框架" in content
test("paper writer zh polish rewrites remaining meta phrases", test_paper_writer_zh_polish_rewrites_remaining_meta_phrases)


def test_paper_writer_resolves_blank_topic_from_project_context():
    from tools.paper_writer import _resolve_effective_topic

    derived_from_project = _resolve_effective_topic(
        "G:/sci",
        "",
        {
            "source_project_path": "G:/demo-project",
            "project_name": "demo-project",
            "topic": "",
        },
    )
    assert derived_from_project == "demo-project"

    derived_from_reference = _resolve_effective_topic(
        "G:/sci",
        "",
        {
            "uploaded_references": [
                {"filename": "smart-grid-dispatch.pdf"},
            ],
        },
    )
    assert derived_from_reference == "smart-grid-dispatch"

    ambiguous_reference_fallback = _resolve_effective_topic(
        "G:/sci",
        "",
        {
            "uploaded_references": [
                {"filename": "paper.pdf"},
                {"filename": "supplementary-material.pdf"},
            ],
        },
    )
    assert ambiguous_reference_fallback == "Research Draft"
test("paper writer derives fallback topic", test_paper_writer_resolves_blank_topic_from_project_context)


def test_paper_writer_budgets_uploaded_reference_context():
    from tools.paper_writer import _build_uploaded_reference_text, _select_uploaded_reference_entries

    uploaded_references = [
        {
            "filename": "robot-navigation-notes.txt",
            "content": (
                "Robot navigation planner latency analysis. "
                "The navigation stack compares controller latency, path smoothness, and lidar-based obstacle handling. "
                * 40
            ),
        },
        {
            "filename": "protein-folding-survey.txt",
            "content": (
                "Protein folding benchmark survey focused on molecular dynamics, residue contact maps, and folding pathways. "
                * 40
            ),
        },
    ]
    project_context = {
        "project_name": "nav-stack",
        "project_summary": "Robot navigation system with planner, controller, lidar mapping, and runtime latency evaluation.",
    }

    selected = _select_uploaded_reference_entries(
        uploaded_references,
        topic="robot navigation latency analysis",
        project_context=project_context,
    )
    assert len(selected) == 1
    assert selected[0]["ref"]["filename"] == "robot-navigation-notes.txt"
    assert len(selected[0]["excerpt"]) <= 900

    injected = _build_uploaded_reference_text(
        uploaded_references,
        topic="robot navigation latency analysis",
        project_context=project_context,
    )
    assert "robot-navigation-notes.txt" in injected
    assert "protein-folding-survey.txt" not in injected
    assert len(injected) <= 2800
test("paper writer budgets uploaded reference context", test_paper_writer_budgets_uploaded_reference_context)


def test_paper_writer_prefers_source_project_outputs_for_evidence_scanning():
    import shutil
    import sys
    import types
    import uuid
    from pathlib import Path

    import tools.paper_writer as paper_writer

    captured: dict[str, object] = {}
    original_runner = sys.modules.get("tools.project_runner")
    original_legacy = paper_writer._legacy_generate_paper_package
    original_backfill = paper_writer._backfill_i18n_and_references

    tmp_root = Path("output") / f"tmp-source-evidence-{uuid.uuid4().hex}"
    workspace_root = tmp_root / "workspace"
    source_root = tmp_root / "source-project"

    (workspace_root / "output" / "figures").mkdir(parents=True)
    (workspace_root / "output" / "results").mkdir(parents=True)
    (source_root / "output" / "figures").mkdir(parents=True)
    (source_root / "output" / "results").mkdir(parents=True)

    (workspace_root / "output" / "figures" / "stale-workspace.png").write_text("stale", encoding="utf-8")
    (workspace_root / "output" / "results" / "stale_metrics.csv").write_text(
        "metric,value\nstale,1\n",
        encoding="utf-8",
    )
    (source_root / "output" / "figures" / "source-proof.png").write_text("source", encoding="utf-8")
    (source_root / "output" / "results" / "source_metrics.csv").write_text(
        "metric,value\nsource,42\n",
        encoding="utf-8",
    )

    fake_runner = types.ModuleType("tools.project_runner")
    fake_runner.run_project = lambda source_path, entry_script=None, timeout=300, dry_run=False, allow_agent_modifications=None: {
        "project_context": {},
        "collected": {"figures": []},
        "run_result": {"success": True},
        "project_type": "test",
    }

    def fake_legacy_generate_paper_package(**kwargs):
        context = dict(kwargs["project_context"])
        captured["project_context"] = context
        return {
            "project_root": str(workspace_root),
            "dashboard_path": str(workspace_root / "output" / "paper-workbench.html"),
            "markdown_path": str(workspace_root / "drafts" / "paper.md"),
            "outline_path": str(workspace_root / "drafts" / "paper-outline.md"),
            "plan_path": str(workspace_root / "drafts" / "paper-plan.md"),
            "prompts_path": str(workspace_root / "drafts" / "paper-prompts.md"),
            "json_path": str(workspace_root / "output" / "paper.json"),
            "latex_path": str(workspace_root / "papers" / "paper.tex"),
            "bib_path": str(workspace_root / "papers" / "references.bib"),
            "artifact": {
                "language": "en",
                "sections": [],
                "references": [],
                "project_context": context,
            },
            "state": {},
        }

    try:
        sys.modules["tools.project_runner"] = fake_runner
        paper_writer._legacy_generate_paper_package = fake_legacy_generate_paper_package
        paper_writer._backfill_i18n_and_references = lambda result: None

        paper_writer.generate_paper_package(
            project_root=workspace_root,
            topic="source evidence test",
            language="en",
            paper_type="general",
            project_context={
                "source_project_path": str(source_root),
                "skip_run": True,
            },
        )
        captured["copied_figure_exists"] = (workspace_root / "output" / "figures" / "source-proof.png").exists()
    finally:
        paper_writer._legacy_generate_paper_package = original_legacy
        paper_writer._backfill_i18n_and_references = original_backfill
        if original_runner is None:
            sys.modules.pop("tools.project_runner", None)
        else:
            sys.modules["tools.project_runner"] = original_runner
        shutil.rmtree(tmp_root, ignore_errors=True)

    project_context = captured["project_context"]
    result_clues = project_context["result_clues"]
    candidate_result_files = project_context["candidate_result_files"]

    assert any("source_metrics" in clue for clue in result_clues)
    assert not any("stale_metrics" in clue for clue in result_clues)
    assert any("output/figures/source-proof.png" in item for item in candidate_result_files)
    assert not any("stale-workspace.png" in item for item in candidate_result_files)
    assert project_context["paper_workspace_figure_files"] == ["source-proof.png"]
    assert captured["copied_figure_exists"] is True
test("paper writer prefers source project outputs for evidence scanning", test_paper_writer_prefers_source_project_outputs_for_evidence_scanning)


def test_writing_refiner_generates_round_outputs():
    import shutil
    import uuid
    from pathlib import Path

    from tools.writing_refiner import refine_document_package

    root = Path("output") / f"tmp-writing-refiner-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        (root / "drafts").mkdir()
        (root / "output").mkdir()
        source = root / "drafts" / "paper-draft.md"
        source.write_text(
            "# 第一章 绪论\n\n"
            "首先，本文围绕移动机器人导航系统展开研究。结果表明，系统具备一定效果。\n\n"
            "```python\nprint('keep')\n```\n\n"
            "需要指出的是，系统在复杂场景下仍然存在定位波动。\n",
            encoding="utf-8",
        )

        round1 = refine_document_package(root, "drafts/paper-draft.md", language="zh")
        round2 = refine_document_package(root, "drafts/paper-draft.md", language="zh")

        assert Path(round1["markdown_path"]).exists()
        assert Path(round2["markdown_path"]).exists()
        assert round1["artifact"]["quality_meta"]["refinement_round"] == 1
        assert round2["artifact"]["quality_meta"]["refinement_round"] == 2

        records_path = root / "output" / "writing-refinement" / "refinement-records.json"
        records = json.loads(records_path.read_text(encoding="utf-8"))
        entry = records["drafts/paper-draft.md"]
        assert len(entry["rounds"]) == 2
        assert entry["rounds"][1]["input_path"] == entry["rounds"][0]["output_path"]
    finally:
        shutil.rmtree(root)
test("writing refiner generates tracked round outputs", test_writing_refiner_generates_round_outputs)


def test_writing_refiner_preserves_markdown_structure():
    import shutil
    import uuid
    from pathlib import Path

    from tools.writing_refiner import refine_document_package

    root = Path("output") / f"tmp-writing-refiner-structure-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        (root / "drafts").mkdir()
        (root / "output").mkdir()
        source = root / "drafts" / "chapter.md"
        source.write_text(
            "# 3. 系统实现\n\n"
            "其次，系统通过参数配置驱动导航流程。\n\n"
            "```yaml\nplanner: teb\ncontroller: mppi\n```\n\n"
            "- 保留这个列表项\n",
            encoding="utf-8",
        )

        result = refine_document_package(root, "drafts/chapter.md", language="zh")
        refined = Path(result["markdown_path"]).read_text(encoding="utf-8")

        assert "# 3. 系统实现" in refined
        assert "```yaml" in refined
        assert "planner: teb" in refined
        assert "- 保留这个列表项" in refined
    finally:
        shutil.rmtree(root)
test("writing refiner preserves markdown structure", test_writing_refiner_preserves_markdown_structure)


def test_crossref():
    from tools.crossref_search import search_by_title
    r = search_by_title("Attention Is All You Need", rows=1)
    assert len(r) > 0 and "attention" in r[0]["title"].lower()
network_test("CrossRef API", test_crossref)

def test_s2():
    from tools.semantic_scholar import search_papers
    r = search_papers("attention is all you need", limit=1)
    assert len(r) > 0
network_test("Semantic Scholar API", test_s2)

def test_arxiv():
    from tools.arxiv_download import search_papers
    r = search_papers("transformer attention mechanism", limit=2)
    assert len(r) > 0
    assert r[0].get("arxiv_id")
network_test("arXiv API", test_arxiv)

def test_verify():
    from tools.unified_search import verify_paper
    r = verify_paper("Attention Is All You Need", authors=["Vaswani"], year=2017)
    assert r.get("verified") == True
    print(f"         (verified_by={r.get('verified_by')})")
network_test("多源交叉验证", test_verify)

def test_auto_search():
    from tools.unified_search import auto_search
    r = auto_search("graph neural network", discipline="cs", limit=3)
    assert len(r) > 0
    sources = set(p.get("_source") for p in r)
    print(f"         (results={len(r)}, sources={sources})")
network_test("统一搜索 (CS学科路由)", test_auto_search)

# ============================================================
def test_writing_profile_guardrails():
    from tools.writing_profiles import build_profile_guardrails

    notes = build_profile_guardrails(
        language="zh",
        role="experiment",
        has_result_evidence=False,
        has_reference_support=False,
    )
    joined = "\n".join(notes)
    assert "阿拉伯数字" in joined
    assert "不伪造精确指标" in joined
    assert "模板句" in joined or "首先" in joined
test("写作约束画像", test_writing_profile_guardrails)


def test_figure_plan_builder():
    from tools.writing_profiles import build_figure_plan, build_figure_plan_summary

    project_context = {
        "project_name": "Demo Planner",
        "project_summary": "一个带实验结果的工程项目",
        "stack": ["Python", "FastAPI"],
        "candidate_source_files": ["src/app.py", "src/pipeline.py"],
        "candidate_config_files": ["config/settings.yaml"],
        "method_clues": ["多阶段处理流程", "推理后处理"],
        "result_clues": ["accuracy: 91.2%, latency: 38ms", "baseline vs method"],
        "candidate_result_files": ["output/results/metrics.csv", "output/figures/compare.png"],
    }

    plan = build_figure_plan(
        topic="项目生成论文",
        language="zh",
        project_context=project_context,
    )
    assert len(plan) >= 3
    assert any(item["figure_type"] == "系统架构图" for item in plan)
    assert any(item["figure_type"] == "结果对比图" for item in plan)
    summary = build_figure_plan_summary(plan, "zh")
    assert "图" in summary
test("图表规划生成", test_figure_plan_builder)


def test_writing_assets_markdown():
    from tools.writing_profiles import render_integrated_writing_assets_markdown

    content = render_integrated_writing_assets_markdown()
    assert "research-writing-skill" in content
    assert "academic-figure-generator" in content
    assert "figure-plan.md" in content
test("写作资产说明", test_writing_assets_markdown)


def test_project_analysis_extracts_structured_assets():
    import os
    import shutil
    import uuid
    from pathlib import Path

    from tools.project_paper_context import analyze_project_for_paper

    root = Path("output") / f"tmp-project-analysis-assets-{uuid.uuid4().hex}"
    project = root / "demo-project"
    (root / "drafts").mkdir(parents=True)
    (root / "output").mkdir(parents=True)
    (project / "src").mkdir(parents=True)
    (project / "results").mkdir(parents=True)
    try:
        (project / "README.md").write_text(
            "# Demo Project\n\nThis project optimizes routing accuracy and latency.\n",
            encoding="utf-8",
        )
        (project / "src" / "model.py").write_text(
            "theta = 0.1\n"
            "alpha = 0.8\n"
            "loss = theta * alpha\n"
            "accuracy = 0.92\n",
            encoding="utf-8",
        )
        (project / "results" / "metrics.csv").write_text(
            "epoch,accuracy,loss\n1,0.81,0.40\n2,0.92,0.25\n",
            encoding="utf-8",
        )
        (project / "results" / "accuracy-curve.png").write_bytes(b"fake-png")
        prior = os.environ.get("SCIPILOT_AGENT_ENABLED")
        os.environ["SCIPILOT_AGENT_ENABLED"] = "0"
        try:
            context = analyze_project_for_paper(root, project, "项目论文测试")
        finally:
            if prior is None:
                os.environ.pop("SCIPILOT_AGENT_ENABLED", None)
            else:
                os.environ["SCIPILOT_AGENT_ENABLED"] = prior

        assert context["figure_candidates"]
        assert context["table_candidates"]
        assert context["equation_candidates"]
        assert context["variable_inventory"]
        assert context["chapter_budget"]["total_figures"] >= 1
        assert context["chapter_budget"]["total_tables"] >= 1
        assert Path(context["analysis_json_path"]).exists()
    finally:
        shutil.rmtree(root)
test("项目分析提取结构化资产", test_project_analysis_extracts_structured_assets)


def test_writing_profile_supports_table_and_equation_plan():
    from tools.writing_profiles import (
        build_equation_plan,
        build_figure_plan,
        build_table_plan,
        render_figure_plan_markdown,
    )

    project_context = {
        "project_name": "Demo Planner",
        "project_summary": "项目包含图、表、公式和指标。",
        "stack": ["Python"],
        "candidate_source_files": ["src/model.py"],
        "candidate_config_files": ["config/settings.yaml"],
        "candidate_result_files": ["output/results/metrics.csv", "output/figures/accuracy-curve.png"],
        "method_clues": ["优化目标", "损失函数"],
        "result_clues": ["accuracy 0.92", "loss 0.25"],
        "figure_candidates": [
            {"caption": "准确率曲线", "path": "output/figures/accuracy-curve.png", "section": "experiment", "role": "experiment"}
        ],
        "table_candidates": [
            {"caption": "主要指标对比", "path": "output/results/metrics.csv", "section": "experiment", "headers": ["指标", "数值"], "rows": [["Accuracy", "0.92"]]}
        ],
        "equation_candidates": [
            {"focus": "损失函数与更新公式", "source": "src/model.py", "section": "theory"}
        ],
        "variable_inventory": [{"symbol": "theta", "evidence": "src/model.py"}],
        "metric_inventory": ["accuracy", "loss"],
        "chapter_budget": {
            "figures": {"experiment": 2},
            "tables": {"experiment": 2},
            "equations": {"theory": 2},
            "total_figures": 1,
            "total_tables": 1,
            "total_equations": 1,
        },
    }

    figure_plan = build_figure_plan(topic="项目论文测试", language="zh", project_context=project_context)
    table_plan = build_table_plan(topic="项目论文测试", language="zh", project_context=project_context)
    equation_plan = build_equation_plan(topic="项目论文测试", language="zh", project_context=project_context)
    markdown = render_figure_plan_markdown(
        topic="项目论文测试",
        language="zh",
        plan=figure_plan,
        table_plan=table_plan,
        equation_plan=equation_plan,
    )

    assert figure_plan
    assert table_plan
    assert equation_plan
    assert "候选表格" in markdown
    assert "候选公式" in markdown
test("图表规划包含表格与公式", test_writing_profile_supports_table_and_equation_plan)


def test_paper_writer_formats_equations_and_relocates_figures():
    from tools.paper_writer import (
        _format_variable_mentions,
        _inject_missing_figure_placeholders,
        _normalize_formula_blocks,
        _number_equation_blocks,
        _relocate_figures_by_chapter,
    )

    md = (
        "## 2. 方法设计\n\n"
        "我们定义 theta 和 loss 作为核心变量。\n\n"
        "$$\n"
        "loss = theta + 1\n"
        "$$\n\n"
        "准确率曲线展示了模型在训练后期的收敛趋势。\n"
    )
    project_context = {
        "variable_inventory": [
            {"symbol": "theta", "evidence": "src/model.py"},
            {"symbol": "loss", "evidence": "src/model.py"},
        ],
        "figure_plan": [
            {
                "caption": "准确率曲线",
                "section": "方法设计",
                "goal": "展示训练收敛趋势",
                "evidence": "accuracy curve",
                "existing_asset": "figs/accuracy.png",
            }
        ],
    }

    formatted = _format_variable_mentions(md, project_context)
    numbered = _number_equation_blocks(formatted, language="zh")
    relocated = _relocate_figures_by_chapter(
        numbered,
        [{"caption": "准确率曲线", "path": "figs/accuracy.png"}],
        language="zh",
        project_context=project_context,
    )

    assert "$theta$" in formatted
    assert "$loss$" in formatted
    assert "\\tag{式2.1}" in numbered
    assert "\n式2.1\n" not in numbered
    assert "![准确率曲线](figs/accuracy.png)" in relocated
    assert relocated.index("![准确率曲线](figs/accuracy.png)") > relocated.index("准确率曲线展示了模型在训练后期的收敛趋势。")
    malformed = (
        "## 2. 方法设计\n\n"
        "$$loss = theta + 1 $$\n"
        "\\tag{2.1}\n"
        "$$\n"
        "式2.1\n"
    )
    repaired = _number_equation_blocks(_normalize_formula_blocks(malformed), language="zh")
    assert repaired.count("\\tag{式2.1}") == 1
    assert "\n式2.1\n" not in repaired
    placeholder = _inject_missing_figure_placeholders(
        "## 2. 方法设计\n\n参数配置关系决定了运行模式切换逻辑。\n",
        {
            "figure_plan": [
                {
                    "caption": "图2 参数配置与运行模式关系",
                    "section": "方法设计",
                    "figure_type": "系统结构图",
                    "goal": "说明参数配置与模式切换的映射关系",
                    "evidence": "config/nav2.yaml",
                }
            ]
        },
        language="zh",
    )
    assert "[此处插入图2-1]" in placeholder
    assert "图2-1 参数配置与运行模式关系" in placeholder
    assert "config/nav2.yaml" not in placeholder
    assert "[待补图]" not in placeholder
test("正文按位置插图并规范公式编号", test_paper_writer_formats_equations_and_relocates_figures)


def test_paper_writer_figure_placeholder_matches_export_protocol():
    from tools.paper_writer import _inject_missing_figure_placeholders
    from tools.research_export import _iter_markdown_blocks

    markdown = _inject_missing_figure_placeholders(
        "## 2. 方法设计\n\n参数配置关系决定了运行模式切换逻辑。\n",
        {
            "figure_plan": [
                {
                    "caption": "图2 参数配置与运行模式关系",
                    "section": "方法设计",
                    "goal": "说明参数配置与模式切换的映射关系",
                }
            ]
        },
        language="zh",
    )

    blocks = _iter_markdown_blocks(markdown.splitlines())
    placeholders = [block for block in blocks if block.get("type") == "image_placeholder"]
    assert placeholders
    assert placeholders[0]["ref"] == "2-1"
    assert placeholders[0].get("caption") in {"", "参数配置与运行模式关系"}
test("paper writer figure placeholder matches export protocol", test_paper_writer_figure_placeholder_matches_export_protocol)


def test_paper_writer_injects_tables_by_plan():
    from tools.paper_writer import _inject_tables_by_plan

    md = (
        "## 4. 实验结果\n\n"
        "主要指标对比展示了方法相对基线的提升。\n\n"
        "进一步分析如下。\n"
    )
    project_context = {
        "table_plan": [
            {
                "caption": "主要指标对比",
                "section": "实验结果",
                "headers": ["指标", "数值"],
                "rows": [["Accuracy", "0.92"], ["Loss", "0.25"]],
            }
        ]
    }

    injected = _inject_tables_by_plan(md, project_context, language="zh")

    assert "表4.1 主要指标对比" in injected
    assert "| 指标 | 数值 |" in injected
    assert injected.index("表4.1 主要指标对比") > injected.index("主要指标对比展示了方法相对基线的提升。")
test("正文按提及位置插入表格", test_paper_writer_injects_tables_by_plan)


def test_paper_writer_repairs_table_headers_and_placeholder_matching():
    from tools.paper_writer import _figure_plan_item_present, _repair_markdown_tables

    table = _repair_markdown_tables(
        "表2.1 planner comparison\n\n"
        "| method | path_length_m | planning_time_ms |\n"
        "| --- | --- | --- |\n"
        "| MPPI | 12.4 | 15.2 |\n"
    )
    assert "| Method | Path Length (m) | Planning Time (ms) |" in table
    assert "| :--- | ---: | ---: |" in table

    existing = (
        "## 2. 方法设计\n\n"
        "![准确率曲线](output/figures/accuracy.png)\n"
        "图2-1 准确率曲线\n"
    )
    assert _figure_plan_item_present(existing, {"caption": "图2 准确率曲线", "existing_asset": "accuracy.png"})
    assert not _figure_plan_item_present(existing, {"caption": "图2 参数配置与运行模式关系", "goal": "说明参数配置与模式切换的映射关系"})
test("表头标准化且补图匹配严格", test_paper_writer_repairs_table_headers_and_placeholder_matching)


def test_paper_writer_keeps_tables_outside_single_line_equations():
    from tools.paper_writer import _normalize_final_manuscript_format

    md = (
        "## 2. Method\n\n"
        "$$loss = x^2$$\n"
        "\u5f0f2.1\n\n"
        "\u88689.9 planner comparison\n"
        "| method | path_length_m | planning_time_ms |\n"
        "| A | 1 | 2 |\n"
    )
    normalized = _normalize_final_manuscript_format(md, language="zh", project_context={})

    assert "\\tag{\u5f0f2.1}" in normalized
    assert "\u88682.1 Planner Comparison" in normalized
    assert "| Method | Path Length (m) | Planning Time (ms) |" in normalized
    assert "planner comparison" not in normalized.split("\\tag{\u5f0f2.1}", 1)[0]


test("单行公式不会吞掉后续表格", test_paper_writer_keeps_tables_outside_single_line_equations)


def test_paper_writer_salvages_broken_formula_fragments():
    from tools.paper_writer import _normalize_final_manuscript_format

    md = (
        "## 3. Design\n\n"
        "系统的基础变量定义如下。$$\n\n"
        "$$x = r \\\\cos(\\\\theta)\n"
        "\\tag{式3.1}\n"
        "$$\n\n"
        "$$\n"
        "$$y = r \\\\sin(\\\\theta)\n"
        "$$式3.2\n"
        "\\tag{式3.2}\n"
        "$$\n\n"
        "z = h\n"
        "\\tag{式3.3}\n"
        "$$\n\n"
        "后续正文继续说明系统实现。\n"
    )
    normalized = _normalize_final_manuscript_format(md, language="zh", project_context={})

    assert "\\tag{式3.1}" in normalized
    assert "\\tag{式3.2}" in normalized
    assert "\\tag{式3.3}" in normalized
    assert "式3.2\n" not in normalized.replace("\\tag{式3.2}\n", "")
    assert "式3.3\n" not in normalized.replace("\\tag{式3.3}\n", "")
    assert "$$\ny = r \\\\sin(\\\\theta)\n\\tag{式3.2}\n$$" in normalized
    assert "$$\nz = h\n\\tag{式3.3}\n$$" in normalized
    assert "后续正文继续说明系统实现。" in normalized


test("坏公式残片会被救援成标准公式块", test_paper_writer_salvages_broken_formula_fragments)


def test_research_export_parses_structured_placeholder_and_repairs_formula_fragments():
    from tools.research_export import _display_equation_tag_text, _extract_formulas_from_text, _iter_markdown_blocks

    blocks = _iter_markdown_blocks(
        [
            "> [此处插入图2-1] 参数配置与运行模式关系",
            "> 图型建议：系统结构图",
            "> 应展示内容：说明参数配置与模式切换的映射关系",
            "> 推荐素材来源：config/nav2.yaml",
            "",
        ]
    )

    assert blocks[0]["type"] == "image_placeholder"
    assert blocks[0]["ref"] == "2-1"
    assert blocks[0]["caption"] == "参数配置与运行模式关系"
    assert blocks[0]["figure_type"] == "系统结构图"
    assert "config/nav2.yaml" in blocks[0]["evidence"]

    segments = _extract_formulas_from_text("$$ R(x,y) = R_0 + 1 \\tag{式2.2}。将式(2.2)代入后续计算。")
    display = [seg for seg in segments if seg["type"] == "display_formula"]
    trailing_text = [seg["content"] for seg in segments if seg["type"] == "text"]

    assert display and display[0]["tag"] == "2.2"
    assert any("将式(2.2)代入后续计算" in content for content in trailing_text)
    assert _display_equation_tag_text("式2.2") == "（式2.2）"


test("research export 识别结构化补图并修复坏公式尾巴", test_research_export_parses_structured_placeholder_and_repairs_formula_fragments)


def test_paper_writer_uploaded_references_stay_secondary_to_project_evidence():
    from tools.paper_writer import _build_uploaded_reference_text, _resolve_effective_topic

    project_context = {
        "project_name": "nav2-stack",
        "source_project_path": "G:/demo/nav2-stack",
        "result_clues": ["planner server", "controller server", "costmap configuration"],
        "uploaded_references": [
            {
                "filename": "finance-report.txt",
                "content": "stock market quarterly earnings guidance and investor summary " * 80,
            },
            {
                "filename": "nav2-controller-notes.txt",
                "content": "planner server controller server costmap configuration and recovery behaviors " * 40,
            },
        ],
    }

    built = _build_uploaded_reference_text(
        project_context["uploaded_references"],
        topic="nav2-stack",
        project_context=project_context,
    )
    resolved = _resolve_effective_topic("G:/sci", "", project_context)

    assert "nav2-controller-notes.txt" in built
    assert "finance-report.txt" not in built
    assert "lower priority than project evidence" in built
    assert resolved == "nav2-stack"


test("uploaded references are supplementary instead of topic-driving", test_paper_writer_uploaded_references_stay_secondary_to_project_evidence)


def test_paper_writer_spreads_figures_across_paragraphs_without_explicit_refs():
    from tools.paper_writer import _relocate_figures_by_chapter

    md = (
        "## 1. Intro\n\n"
        "\u7b2c\u4e00\u6bb5\u8bf4\u660e\u7cfb\u7edf\u7ed3\u6784\u3002\n\n"
        "\u7b2c\u4e8c\u6bb5\u8bf4\u660e\u53c2\u6570\u914d\u7f6e\u3002\n\n"
        "\u7b2c\u4e09\u6bb5\u8bf4\u660e\u63a7\u5236\u6d41\u7a0b\u3002\n"
    )
    relocated = _relocate_figures_by_chapter(
        md,
        [
            {"caption": "\u7ed3\u6784\u6846\u56fe", "path": "a.png"},
            {"caption": "\u53c2\u6570\u5173\u7cfb\u56fe", "path": "b.png"},
            {"caption": "\u63a7\u5236\u6d41\u7a0b\u56fe", "path": "c.png"},
        ],
        language="zh",
        project_context={},
    )

    assert relocated.index("![\u7ed3\u6784\u6846\u56fe](a.png)") < relocated.index("![\u53c2\u6570\u5173\u7cfb\u56fe](b.png)")
    assert relocated.index("![\u53c2\u6570\u5173\u7cfb\u56fe](b.png)") < relocated.index("![\u63a7\u5236\u6d41\u7a0b\u56fe](c.png)")
    assert relocated.index("![\u7ed3\u6784\u6846\u56fe](a.png)") < relocated.index("\u7b2c\u4e8c\u6bb5\u8bf4\u660e\u53c2\u6570\u914d\u7f6e\u3002")
    assert relocated.index("![\u53c2\u6570\u5173\u7cfb\u56fe](b.png)") < relocated.index("\u7b2c\u4e09\u6bb5\u8bf4\u660e\u63a7\u5236\u6d41\u7a0b\u3002")


test("无显式图引用时仍按段落分散插图", test_paper_writer_spreads_figures_across_paragraphs_without_explicit_refs)


def test_research_export_parses_blockquote_figure_placeholders_without_blank_lines():
    from tools.research_export import _iter_markdown_blocks

    lines = [
        "\u4e0a\u4e00\u6bb5\u8bf4\u660e\u9644\u8fd1\u9700\u8981\u4e00\u5f20\u56fe\u3002",
        "> [\u5f85\u8865\u56fe] \u8def\u5f84\u89c4\u5212\u6d41\u7a0b\u56fe",
        "> \u56fe\u578b\u5efa\u8bae\uff1a\u6d41\u7a0b\u56fe",
        "> \u5e94\u5c55\u793a\u5185\u5bb9\uff1a\u5c55\u793a\u5b9a\u4f4d\u3001\u5efa\u56fe\u4e0e\u8def\u5f84\u89c4\u5212\u5173\u7cfb",
        "> \u63a8\u8350\u7d20\u6750\u6765\u6e90\uff1aoutput/figures/planner-placeholder.png",
        "\u540e\u7eed\u6bb5\u843d\u7ee7\u7eed\u89e3\u91ca\u3002",
    ]

    blocks = _iter_markdown_blocks(lines)
    placeholder = next(block for block in blocks if block.get("type") == "image_placeholder")

    assert placeholder["caption"] == "\u8def\u5f84\u89c4\u5212\u6d41\u7a0b\u56fe"
    assert placeholder["figure_type"] == "\u6d41\u7a0b\u56fe"
    assert "\u5b9a\u4f4d" in placeholder["goal"]
    assert placeholder["evidence"] == "output/figures/planner-placeholder.png"
    assert all("[\u5f85\u8865\u56fe]" not in block.get("text", "") for block in blocks if block.get("type") == "paragraph")


test("research export parses quoted figure placeholders", test_research_export_parses_blockquote_figure_placeholders_without_blank_lines)


def test_research_export_docx_salvages_formula_and_embeds_quoted_placeholder_image():
    import base64
    import shutil
    import uuid
    from pathlib import Path

    from docx import Document as DocxDocument
    from tools.research_export import export_markdown_to_docx

    png_bytes = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO7Z0L8AAAAASUVORK5CYII="
    )

    root = Path("output") / f"tmp-research-export-{uuid.uuid4().hex}"
    root.mkdir(parents=True)
    try:
        (root / "drafts").mkdir()
        (root / "output" / "figures").mkdir(parents=True)

        image_path = root / "output" / "figures" / "planner-placeholder.png"
        image_path.write_bytes(png_bytes)

        markdown = (
            "## 2. Method\n\n"
            "The nearby paragraph expects a figure.\n"
            "> [\u5f85\u8865\u56fe] \u8def\u5f84\u89c4\u5212\u6d41\u7a0b\u56fe\n"
            "> \u56fe\u578b\u5efa\u8bae\uff1a\u6d41\u7a0b\u56fe\n"
            "> \u5e94\u5c55\u793a\u5185\u5bb9\uff1a\u5c55\u793a\u5b9a\u4f4d\u3001\u5efa\u56fe\u4e0e\u8def\u5f84\u89c4\u5212\u5173\u7cfb\n"
            f"> \u63a8\u8350\u7d20\u6750\u6765\u6e90\uff1a{(Path('output') / 'figures' / image_path.name).as_posix()}\n"
            "Further discussion continues here.\n\n"
            "$$ f(n)=g(n)+h(n) \\tag{2.2}\u3002\u5c06\u5f0f(2.2)\u4ee3\u5165\u603b\u4ee3\u4ef7\u51fd\u6570\n"
            "\\tag{\u5f0f2.2} $$\n"
        )
        (root / "drafts" / "paper-draft.md").write_text(markdown, encoding="utf-8")

        result = export_markdown_to_docx(project_root=root, artifact="paper")
        doc = DocxDocument(result["output_path"])
        paragraph_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)

        assert len(doc.inline_shapes) == 1
        assert "\u8def\u5f84\u89c4\u5212\u6d41\u7a0b\u56fe" in paragraph_text
        assert "\u5c06\u5f0f(2.2)\u4ee3\u5165\u603b\u4ee3\u4ef7\u51fd\u6570" in paragraph_text
        assert "$$" not in paragraph_text
        assert "\\tag{" not in paragraph_text
    finally:
        shutil.rmtree(root)


test("research export docx salvages formula and embeds quoted placeholder image", test_research_export_docx_salvages_formula_and_embeds_quoted_placeholder_image)


print(f"\n{'='*50}")
print(f"  PASS: {PASS}  |  FAIL: {FAIL}  |  SKIP: {SKIP}")
print(f"{'='*50}")
if FAIL > 0:
    sys.exit(1)
